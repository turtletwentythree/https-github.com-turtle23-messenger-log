#!/usr/bin/env python3
"""
Sync the Turtle23 Messenger booking queue from SharePoint and rebuild the
static web page (docs/index.html) that GitHub Pages serves.

Flow:
  1. Get an app-only Microsoft Graph access token via the client-credentials
     flow (Azure AD app registration; no user sign-in involved).
  2. Resolve the SharePoint sharing URL to a driveItem via the /shares
     endpoint, and download the workbook's raw bytes.
  3. Parse "Sheet1" with openpyxl, exactly mirroring the column layout of
     the Microsoft-Forms-backed export (same indices verified against the
     manually-exported CSV during development).
  4. Write scripts/../docs/booking_data.json (compact records) and rebuild
     docs/index.html from scripts/messenger_log.tpl.html.
  5. Build today's (Asia/Bangkok) print-ready job-sheet workbook (4 sheets,
     see job_sheet.py) and upload it back to the SAME SharePoint folder as
     the original template file, named "MESSENGER {d}-{m}-{yyyy}.xlsx".
     Re-running later the same day overwrites that day's file in place, so
     it stays current as new bookings come in.

Required environment variables (set as GitHub Actions secrets):
  AZURE_TENANT_ID       - Azure AD tenant ID (GUID or verified domain)
  AZURE_CLIENT_ID       - App registration's Application (client) ID
  AZURE_CLIENT_SECRET   - App registration's client secret VALUE
  SHAREPOINT_FILE_URL   - The full sharing URL to the queue .xlsx file, e.g.:
                          https://turtle23-my.sharepoint.com/:x:/r/personal/
                          sarocha_k_turtle23_com/_layouts/15/Doc.aspx?sourcedoc=...
  SHAREPOINT_TEMPLATE_URL - The full sharing URL to "MESSENGER From 1 2 (2).xlsx"
                          (used only to locate its parent folder — the daily
                          job sheet is uploaded next to it, never overwriting it)

Optional:
  SHEET_NAME             - defaults to "Sheet1"
  SKIP_JOB_SHEET_UPLOAD  - set to "1" to only sync docs/ and skip step 5

Required Graph API application permissions (admin-consented):
  Sites.Read.All (or Files.Read.All) is NOT enough on its own once uploading
  is enabled — you also need write access: Sites.ReadWrite.All or
  Files.ReadWrite.All.
"""
import base64
import datetime
import json
import os
import sys
from io import BytesIO

import requests
import openpyxl

import job_sheet

GRAPH_BASE = "https://graph.microsoft.com/v1.0"

# --- Column layout (0-indexed), verified against the Microsoft-Forms export ---
COL_ID = 0
COL_START_TIME = 1          # "Start time" -> used as the "submitted" timestamp
COL_DATE = 6
COL_DEPT = 7
COL_REQUESTER = 8
COL_REQUESTER_PHONE = 9
COL_TASK_TYPE = 10
COL_LOCATION = 12
COL_LOCATION_FALLBACKS = [13, 30, 43, 14]
COL_BANK_NAME = 41
COL_BANK_DETAIL = 42
COL_DEST_CONTACT = 45
COL_DEST_PHONE = 46
COL_NOTES = 47
COL_SHIFT = 48


def get_access_token():
    tenant_id = os.environ["AZURE_TENANT_ID"]
    client_id = os.environ["AZURE_CLIENT_ID"]
    client_secret = os.environ["AZURE_CLIENT_SECRET"]

    url = f"https://login.microsoftonline.com/{tenant_id}/oauth2/v2.0/token"
    resp = requests.post(
        url,
        data={
            "client_id": client_id,
            "client_secret": client_secret,
            "scope": "https://graph.microsoft.com/.default",
            "grant_type": "client_credentials",
        },
        timeout=30,
    )
    if not resp.ok:
        print("Failed to obtain Graph access token:", resp.status_code, resp.text, file=sys.stderr)
        resp.raise_for_status()
    return resp.json()["access_token"]


def encode_sharing_url(url):
    """Encode a SharePoint sharing URL into the Graph 'shares' id format."""
    b64 = base64.b64encode(url.encode("utf-8")).decode("utf-8")
    encoded = b64.replace("/", "_").replace("+", "-").rstrip("=")
    return "u!" + encoded


def resolve_share_item(token, sharing_url):
    """Resolve a SharePoint sharing URL to its driveItem (includes parentReference)."""
    share_id = encode_sharing_url(sharing_url)
    headers = {"Authorization": f"Bearer {token}"}
    # ?$expand=parentReference isn't needed — parentReference is returned by default.
    resp = requests.get(f"{GRAPH_BASE}/shares/{share_id}/driveItem", headers=headers, timeout=30)
    if not resp.ok:
        print("Failed to resolve SharePoint share to a driveItem:", resp.status_code, resp.text, file=sys.stderr)
        resp.raise_for_status()
    return resp.json()


def download_workbook_bytes(token, sharing_url):
    item = resolve_share_item(token, sharing_url)
    print(f"Resolved driveItem: {item.get('name')} ({item.get('id')}), size={item.get('size')} bytes")

    share_id = encode_sharing_url(sharing_url)
    headers = {"Authorization": f"Bearer {token}"}
    resp = requests.get(f"{GRAPH_BASE}/shares/{share_id}/driveItem/content", headers=headers, timeout=120)
    if not resp.ok:
        print("Failed to download workbook content:", resp.status_code, resp.text, file=sys.stderr)
        resp.raise_for_status()
    return resp.content


def upload_job_sheet(token, template_sharing_url, filename, content_bytes):
    """Upload content_bytes as `filename` into the same folder as the file
    pointed to by template_sharing_url (never touching the template itself)."""
    template_item = resolve_share_item(token, template_sharing_url)
    parent = template_item.get("parentReference") or {}
    drive_id = parent.get("driveId")
    parent_path = parent.get("path", "")  # e.g. "/drives/{driveId}/root:/Messenger-SKOOTAR"
    if not drive_id:
        raise RuntimeError(f"Could not determine target drive/folder from template item: {template_item}")

    marker = "root:"
    idx = parent_path.find(marker)
    folder_path = parent_path[idx + len(marker):] if idx != -1 else ""  # e.g. "/Messenger-SKOOTAR"

    headers = {
        "Authorization": f"Bearer {token}",
        "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    }
    url = f"{GRAPH_BASE}/drives/{drive_id}/root:{folder_path}/{filename}:/content"
    resp = requests.put(url, headers=headers, data=content_bytes, timeout=120)
    if not resp.ok:
        print("Failed to upload job sheet to SharePoint:", resp.status_code, resp.text, file=sys.stderr)
        resp.raise_for_status()
    uploaded = resp.json()
    print(f"Uploaded job sheet: {uploaded.get('webUrl')}")
    return uploaded


def resolve_location(row):
    loc = (row[COL_LOCATION] or "").strip()
    if loc and loc != "อื่นๆ":
        return loc
    if loc == "อื่นๆ":
        other = (row[43] or "").strip()
        if other:
            return other
    for i in COL_LOCATION_FALLBACKS:
        v = (row[i] or "").strip()
        if v and v != "อื่นๆ":
            return v
    # NOTE: intentionally NOT falling back to the bank name/detail columns
    # here — this matches the web app's existing booking_data.json, which
    # leaves these ~3 edge-case rows blank rather than synthesizing a
    # "ธนาคาร ..." location string (that synthesis is only used by the
    # Excel job-sheet generator, a separate code path).
    return loc


def normalize_shift(s):
    s = (s or "").strip()
    if "เช้า" in s:
        return "เช้า"
    if "บ่าย" in s:
        return "บ่าย"
    return ""


def to_iso_date(s):
    s = (s or "").strip()
    if not s:
        return ""
    try:
        d, m, y = s.split("/")
        return f"{int(y):04d}-{int(m):02d}-{int(d):02d}"
    except ValueError:
        return ""


def cell_str(v):
    if v is None:
        return ""
    return str(v).strip()


def load_records(xlsx_bytes, sheet_name):
    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes), data_only=True, read_only=True)
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.worksheets[0]

    records = []
    rows_iter = ws.iter_rows(values_only=True)
    next(rows_iter, None)  # skip header row

    for raw in rows_iter:
        if raw is None or len(raw) < 49:
            continue
        row = [cell_str(v) for v in raw]
        if not row[COL_ID].isdigit():
            continue
        records.append({
            "id": row[COL_ID],
            "date": to_iso_date(row[COL_DATE]),
            "dept": row[COL_DEPT],
            "requester": row[COL_REQUESTER],
            "phone": row[COL_REQUESTER_PHONE],
            "task": row[COL_TASK_TYPE],
            "loc": resolve_location(row),
            "contact": row[COL_DEST_CONTACT],
            "contact_phone": row[COL_DEST_PHONE],
            "notes": row[COL_NOTES],
            "shift": normalize_shift(row[COL_SHIFT]),
            "submitted": row[COL_START_TIME],
        })
    return records


def bangkok_now():
    return datetime.datetime.now(datetime.timezone.utc).astimezone(
        datetime.timezone(datetime.timedelta(hours=7))  # Asia/Bangkok
    )


def build_html(records, repo_root, now):
    tpl_path = os.path.join(repo_root, "scripts", "messenger_log.tpl.html")
    with open(tpl_path, encoding="utf-8") as f:
        tpl = f.read()

    generated_at = now.strftime("%d/%m/%Y %H:%M น.")
    today = now.strftime("%Y-%m-%d")

    data_json = json.dumps(records, ensure_ascii=False, separators=(",", ":"))
    html = (
        tpl.replace("__BOOKING_DATA_JSON__", data_json)
        .replace("__GENERATED_AT__", generated_at)
        .replace("__TODAY__", today)
    )

    docs_dir = os.path.join(repo_root, "docs")
    os.makedirs(docs_dir, exist_ok=True)
    with open(os.path.join(docs_dir, "index.html"), "w", encoding="utf-8") as f:
        f.write(html)
    with open(os.path.join(docs_dir, "booking_data.json"), "w", encoding="utf-8") as f:
        json.dump(records, f, ensure_ascii=False, indent=1)

    return generated_at


def main():
    repo_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    sharing_url = os.environ["SHAREPOINT_FILE_URL"]
    template_sharing_url = os.environ.get("SHAREPOINT_TEMPLATE_URL")
    sheet_name = os.environ.get("SHEET_NAME", "Sheet1")
    skip_upload = os.environ.get("SKIP_JOB_SHEET_UPLOAD") == "1"

    token = get_access_token()
    xlsx_bytes = download_workbook_bytes(token, sharing_url)
    records = load_records(xlsx_bytes, sheet_name)
    if not records:
        print("WARNING: parsed 0 records — refusing to overwrite docs/ with empty data.", file=sys.stderr)
        sys.exit(1)

    now = bangkok_now()
    generated_at = build_html(records, repo_root, now)
    print(f"Synced {len(records)} records. Page generated at {generated_at}.")

    if skip_upload:
        print("SKIP_JOB_SHEET_UPLOAD=1 — not building/uploading the daily job sheet.")
        return
    if not template_sharing_url:
        print("SHAREPOINT_TEMPLATE_URL not set — not building/uploading the daily job sheet.")
        return

    job_bytes, job_filename = job_sheet.build_workbook(records, now.date())
    upload_job_sheet(token, template_sharing_url, job_filename, job_bytes)
    print(f"Uploaded today's job sheet as {job_filename}.")


if __name__ == "__main__":
    main()
